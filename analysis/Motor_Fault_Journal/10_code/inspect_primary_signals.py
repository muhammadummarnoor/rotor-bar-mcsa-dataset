from pathlib import Path
from collections import defaultdict
import csv, math, statistics
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1] / "04_raw_data"

def numeric(s):
    if s is None: return None
    s=str(s).strip()
    if "->" in s: s=s.split("->")[-1].strip()
    try: return float(s)
    except: return None

def summarize(name, cols):
    print("\n", name)
    for j,vals in sorted(cols.items()):
        clean=[x for x in vals if x is not None and math.isfinite(x)]
        if not clean: continue
        s=sorted(clean); n=len(s)
        q=lambda p:s[min(n-1,int((n-1)*p))]
        lows=sum(x<100 for x in clean); mids=sum(100<=x<500 for x in clean); highs=sum(x>=500 for x in clean)
        transitions=[]; prev=clean[0]<100
        for i,x in enumerate(clean[1:],2):
            now=x<100
            if now!=prev and len(transitions)<12: transitions.append(i)
            prev=now
        print(f"col{j+1}: n={n} min={s[0]:.3g} q01={q(.01):.3g} median={q(.5):.3g} q99={q(.99):.3g} max={s[-1]:.3g} <100={lows} 100-499={mids} >=500={highs} transitions={transitions}")

for p in sorted(root.rglob("*")):
    if not p.is_file() or p.suffix.lower() not in {".csv",".xlsx"}: continue
    cols=defaultdict(list)
    if p.suffix.lower()==".csv":
        with p.open('r',encoding='utf-8-sig',errors='replace',newline='') as f:
            for row in csv.reader(f):
                for j,x in enumerate(row): cols[j].append(numeric(x))
    else:
        wb=load_workbook(p,read_only=True,data_only=True)
        for row in wb.active.iter_rows(values_only=True):
            for j,x in enumerate(row): cols[j].append(numeric(x))
    summarize(str(p.relative_to(root)),cols)
