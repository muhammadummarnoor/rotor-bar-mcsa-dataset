from __future__ import annotations

from pathlib import Path
from collections import Counter
import csv
import hashlib
import json
import math
import time

import joblib
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl import load_workbook
from scipy import signal, stats
from sklearn.base import clone
from sklearn.compose import TransformedTargetRegressor
from sklearn.dummy import DummyClassifier
from sklearn.ensemble import (
    ExtraTreesClassifier,
    HistGradientBoostingClassifier,
    RandomForestClassifier,
)
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    classification_report,
    confusion_matrix,
    f1_score,
    matthews_corrcoef,
    precision_score,
    recall_score,
)
from sklearn.neighbors import KNeighborsClassifier
from sklearn.naive_bayes import GaussianNB
from sklearn.neural_network import MLPClassifier
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVC
from sklearn.tree import DecisionTreeClassifier

PROJECT = Path(__file__).resolve().parents[1]
RAW = PROJECT / "04_raw_data"
MANIFEST = PROJECT / "05_data_manifest"
PROCESSED = PROJECT / "06_processed_data"
FEATURES = PROJECT / "07_features"
MODELS = PROJECT / "08_models"
RESULTS = PROJECT / "09_results"

FS = 650.0
WINDOW = 650
GUARD = 650
MIN_SIGNAL_RUN = 5_000
BLOCK = 9_000
RANDOM_STATE = 20260824
CLASS_ORDER = ["healthy", "drb_1", "drb_2", "drb_3"]

SOURCE_MAP = {
    "healthy": RAW / "healthy" / "healthy_2022-02-17.csv",
    "drb_1": RAW / "drb_1" / "drb1_2022-03-04.xlsx",
    "drb_2": RAW / "drb_2" / "drb2_2022-03-04.csv",
    "drb_3": RAW / "drb_3" / "drb3_2022-03-04.csv",
}


def ensure_dirs():
    for p in [MANIFEST, PROCESSED, FEATURES, MODELS, RESULTS / "tables",
              RESULTS / "confusion_matrices", RESULTS / "feature_importance",
              RESULTS / "validation_checks"]:
        p.mkdir(parents=True, exist_ok=True)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def numeric(x):
    if x is None:
        return np.nan
    s = str(x).strip()
    if "->" in s:
        s = s.split("->")[-1].strip()
    try:
        return float(s)
    except Exception:
        return np.nan


def read_columns(path: Path) -> list[np.ndarray]:
    cols: dict[int, list[float]] = {}
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
            for row in csv.reader(f):
                for j, value in enumerate(row):
                    cols.setdefault(j, []).append(numeric(value))
                for j in range(len(row), len(cols)):
                    cols[j].append(np.nan)
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        for row in wb.active.iter_rows(values_only=True):
            for j, value in enumerate(row):
                cols.setdefault(j, []).append(numeric(value))
    return [np.asarray(cols[j], dtype=float) for j in sorted(cols)]


def contiguous_runs(values: np.ndarray, predicate, minimum: int = MIN_SIGNAL_RUN):
    mask = np.isfinite(values) & predicate(values)
    edges = np.diff(np.r_[False, mask, False].astype(int))
    starts = np.where(edges == 1)[0]
    ends = np.where(edges == -1)[0]
    return [(int(s), int(e), values[s:e].copy()) for s, e in zip(starts, ends) if e - s >= minimum]


def extract_current_runs():
    records = []
    manifest_rows = []
    for label, path in SOURCE_MAP.items():
        columns = read_columns(path)
        for col_idx, col in enumerate(columns):
            for run_idx, (start, end, values) in enumerate(
                contiguous_runs(col, lambda x: (x >= 500) & (x <= 1100)), 1
            ):
                run_id = f"{label}_c{col_idx+1}_r{run_idx}"
                records.append(dict(label=label, run_id=run_id, values=values,
                                    source=path, column=col_idx + 1,
                                    start_row=start + 1, end_row=end))
                manifest_rows.append(dict(
                    class_label=label,
                    source_file=str(path.relative_to(PROJECT)),
                    source_sha256=sha256(path),
                    source_column=col_idx + 1,
                    run_id=run_id,
                    start_row=start + 1,
                    end_row=end,
                    samples=len(values),
                    inferred_signal="current_adc_like",
                    label_basis="source filename plus rotor photographs",
                    independence_status="contiguous block; physical restart not confirmed",
                ))
    pd.DataFrame(manifest_rows).to_csv(MANIFEST / "current_run_manifest.csv", index=False)
    return records


def make_guarded_blocks(records):
    """Create contiguous pseudo-recording blocks and assign whole blocks to splits.

    This prevents window overlap across splits but does not convert one acquisition
    into independent physical sessions.
    """
    by_class = {c: [] for c in CLASS_ORDER}
    block_manifest = []
    for rec in records:
        x = rec["values"]
        cursor = 0
        block_no = 0
        while cursor + BLOCK <= len(x):
            block_no += 1
            vals = x[cursor:cursor + BLOCK].copy()
            block_id = f'{rec["run_id"]}_b{block_no}'
            block_record = dict(rec)
            block_record.update(block_id=block_id, values=vals,
                                block_start=cursor, block_end=cursor + BLOCK)
            by_class[rec["label"]].append(block_record)
            cursor += BLOCK + GUARD

    # Deterministic round-robin allocation within each class. Entire blocks stay intact.
    # Every class must contribute at least one block to each split.
    split_map = {}
    for label, blocks in by_class.items():
        if len(blocks) < 3:
            raise RuntimeError(f"Class {label} has only {len(blocks)} usable blocks; need at least 3")
        for i, b in enumerate(blocks):
            split = ("train", "validation", "test")[i % 3]
            split_map[b["block_id"]] = split
            block_manifest.append(dict(
                class_label=label,
                run_id=b["run_id"],
                block_id=b["block_id"],
                split=split,
                source_file=str(b["source"].relative_to(PROJECT)),
                source_column=b["column"],
                source_start_row=b["start_row"],
                source_end_row=b["end_row"],
                within_run_block_start=b["block_start"],
                within_run_block_end=b["block_end"],
                guard_samples_after=GUARD,
            ))
    pd.DataFrame(block_manifest).to_csv(MANIFEST / "split_block_manifest.csv", index=False)
    return by_class, split_map


def spectral_band_power(freq, power, low, high):
    m = (freq >= low) & (freq < high)
    return float(power[m].sum()) if m.any() else 0.0


def features_for_window(x: np.ndarray) -> dict[str, float]:
    x = np.asarray(x, dtype=float)
    x = x - np.mean(x)
    rms = float(np.sqrt(np.mean(x * x)))
    abs_mean = float(np.mean(np.abs(x)))
    peak = float(np.max(np.abs(x)))
    p2p = float(np.ptp(x))
    std = float(np.std(x, ddof=1))
    energy = float(np.mean(x * x))
    crest = peak / rms if rms else 0.0
    shape = rms / abs_mean if abs_mean else 0.0
    impulse = peak / abs_mean if abs_mean else 0.0
    root_abs = float(np.mean(np.sqrt(np.abs(x))))
    clearance = peak / (root_abs * root_abs) if root_abs else 0.0
    win = np.hanning(len(x))
    fft = np.fft.rfft(x * win)
    power = np.abs(fft) ** 2
    freq = np.fft.rfftfreq(len(x), d=1.0 / FS)
    power_sum = float(power.sum())
    prob = power / power_sum if power_sum else np.zeros_like(power)
    nz = prob > 0
    entropy = float(-(prob[nz] * np.log2(prob[nz])).sum()) if nz.any() else 0.0
    centroid = float((freq * power).sum() / power_sum) if power_sum else 0.0
    spread = float(np.sqrt((((freq - centroid) ** 2) * power).sum() / power_sum)) if power_sum else 0.0
    peak_idx = int(np.argmax(power[1:]) + 1) if len(power) > 1 else 0
    fundamental = spectral_band_power(freq, power, 45.0, 55.0)
    total_non_dc = float(power[1:].sum()) if len(power) > 1 else 0.0
    return {
        "mean_centered": float(np.mean(x)),
        "std": std,
        "rms": rms,
        "peak_abs": peak,
        "peak_to_peak": p2p,
        "skewness": float(stats.skew(x, bias=False, nan_policy="omit")),
        "kurtosis": float(stats.kurtosis(x, fisher=True, bias=False, nan_policy="omit")),
        "crest_factor": crest,
        "shape_factor": shape,
        "impulse_factor": impulse,
        "clearance_factor": clearance,
        "median_abs_deviation": float(stats.median_abs_deviation(x)),
        "interquartile_range": float(stats.iqr(x)),
        "energy": energy,
        "spectral_entropy": entropy,
        "spectral_centroid_hz": centroid,
        "spectral_spread_hz": spread,
        "dominant_frequency_hz": float(freq[peak_idx]),
        "bandpower_0_25": spectral_band_power(freq, power, 0.0, 25.0),
        "bandpower_25_45": spectral_band_power(freq, power, 25.0, 45.0),
        "bandpower_45_55": fundamental,
        "bandpower_55_100": spectral_band_power(freq, power, 55.0, 100.0),
        "bandpower_100_325": spectral_band_power(freq, power, 100.0, FS / 2 + 1e-9),
        "fundamental_power_ratio": fundamental / total_non_dc if total_non_dc else 0.0,
    }


def build_feature_table(by_class, split_map):
    rows = []
    hashes = set()
    for label, blocks in by_class.items():
        for b in blocks:
            split = split_map[b["block_id"]]
            vals = b["values"]
            for start in range(0, len(vals) - WINDOW + 1, WINDOW):
                w = vals[start:start + WINDOW]
                window_hash = hashlib.sha256(np.asarray(w, dtype=np.float64).tobytes()).hexdigest()
                if window_hash in hashes:
                    # Duplicate windows are excluded rather than allowed to inflate results.
                    continue
                hashes.add(window_hash)
                row = features_for_window(w)
                row.update(dict(label=label, split=split, run_id=b["run_id"],
                                block_id=b["block_id"], window_start=start,
                                window_hash=window_hash))
                rows.append(row)
    df = pd.DataFrame(rows)
    # mean_centered is theoretically zero after per-window centering.  Its tiny
    # floating-point residue is not a physical feature and is excluded from all
    # final models while retained in the extracted table for auditability.
    feature_cols = [c for c in df.columns if c not in
                    {"label", "split", "run_id", "block_id", "window_start", "window_hash",
                     "mean_centered"}]
    df.to_csv(FEATURES / "current_window_features.csv", index=False)

    # Leakage audit.
    checks = {
        "rows": len(df),
        "class_by_split": df.groupby(["split", "label"]).size().unstack(fill_value=0).to_dict(),
        "unique_window_hashes": int(df["window_hash"].nunique()),
        "duplicate_window_hashes": int(df["window_hash"].duplicated().sum()),
        "block_cross_split_count": int((df.groupby("block_id")["split"].nunique() > 1).sum()),
        "window_length_samples": WINDOW,
        "sampling_rate_hz_assumed": FS,
        "guard_samples": GUARD,
        "limitation": "Blocks are non-overlapping and guard-separated, but some originate from the same acquisition session.",
    }
    (RESULTS / "validation_checks" / "leakage_audit.json").write_text(
        json.dumps(checks, indent=2), encoding="utf-8"
    )
    return df, feature_cols


def candidates():
    scaled = lambda model: Pipeline([
        ("imputer", SimpleImputer(strategy="median")),
        ("scaler", StandardScaler()),
        ("model", model),
    ])
    unscaled = lambda model: Pipeline([
        ("imputer", SimpleImputer(strategy="median")),
        ("model", model),
    ])
    return {
        "Dummy": [unscaled(DummyClassifier(strategy="most_frequent"))],
        "LogisticRegression": [scaled(LogisticRegression(C=c, max_iter=5000, class_weight="balanced", random_state=RANDOM_STATE)) for c in [0.1, 1.0, 10.0]],
        "GaussianNB": [scaled(GaussianNB(var_smoothing=v)) for v in [1e-11, 1e-9, 1e-7]],
        "KNN": [scaled(KNeighborsClassifier(n_neighbors=k, weights=w)) for k in [3, 5, 9] for w in ["uniform", "distance"]],
        # Probability estimation remains disabled by SVC's default. Omitting
        # the deprecated keyword keeps scikit-learn 1.9+ runs warning-free.
        "LinearSVM": [scaled(SVC(kernel="linear", C=c, class_weight="balanced", random_state=RANDOM_STATE)) for c in [0.1, 1.0, 10.0]],
        "RBFSVM": [scaled(SVC(kernel="rbf", C=c, gamma=g, class_weight="balanced", random_state=RANDOM_STATE)) for c in [1.0, 10.0, 100.0] for g in ["scale", 0.01]],
        "DecisionTree": [unscaled(DecisionTreeClassifier(max_depth=d, min_samples_leaf=l, class_weight="balanced", random_state=RANDOM_STATE)) for d in [3, 6, None] for l in [1, 3]],
        "RandomForest": [unscaled(RandomForestClassifier(n_estimators=400, max_depth=d, min_samples_leaf=l, class_weight="balanced", n_jobs=1, random_state=RANDOM_STATE)) for d in [6, None] for l in [1, 3]],
        "ExtraTrees": [unscaled(ExtraTreesClassifier(n_estimators=400, max_depth=d, min_samples_leaf=l, class_weight="balanced", n_jobs=1, random_state=RANDOM_STATE)) for d in [6, None] for l in [1, 3]],
        "HistGradientBoosting": [unscaled(HistGradientBoostingClassifier(learning_rate=lr, max_leaf_nodes=n, l2_regularization=reg, random_state=RANDOM_STATE)) for lr in [0.05, 0.1] for n in [7, 15] for reg in [0.0, 1.0]],
        "MLP": [scaled(MLPClassifier(hidden_layer_sizes=h, alpha=a, early_stopping=True, max_iter=2000, random_state=RANDOM_STATE)) for h in [(32,), (64, 32)] for a in [0.0001, 0.01]],
    }


def evaluate(df, feature_cols):
    split = {s: df[df["split"] == s].copy() for s in ["train", "validation", "test"]}
    X = {s: split[s][feature_cols] for s in split}
    y = {s: split[s]["label"] for s in split}
    rows = []
    fitted = {}
    for family, variants in candidates().items():
        for idx, model in enumerate(variants):
            t0 = time.perf_counter()
            model.fit(X["train"], y["train"])
            fit_seconds = time.perf_counter() - t0
            pred_val = model.predict(X["validation"])
            fitted[(family, idx)] = model
            rows.append(dict(
                classifier=family,
                variant_index=idx,
                validation_macro_f1=f1_score(y["validation"], pred_val, average="macro", zero_division=0),
                validation_balanced_accuracy=balanced_accuracy_score(y["validation"], pred_val),
                validation_mcc=matthews_corrcoef(y["validation"], pred_val),
                training_seconds=fit_seconds,
            ))

    # The test set is not touched during candidate ranking.  Candidate order is
    # predeclared from simpler to more complex settings and is the final tie
    # breaker after validation metrics and classifier name.
    leaderboard = pd.DataFrame(rows).sort_values(
        ["validation_macro_f1", "validation_balanced_accuracy", "validation_mcc", "classifier", "variant_index"],
        ascending=[False, False, False, True, True],
        kind="mergesort",
    ).reset_index(drop=True)
    leaderboard.to_csv(RESULTS / "tables" / "validation_candidate_leaderboard.csv", index=False)
    selected = leaderboard.iloc[0]
    winner = selected["classifier"]
    winner_idx = int(selected["variant_index"])

    # After configuration selection is frozen, refit a fresh clone using all
    # development data.  Test labels and features remain untouched until the
    # single prediction call below.
    development = pd.concat([split["train"], split["validation"]], ignore_index=True)
    final_model = clone(candidates()[winner][winner_idx])
    t0 = time.perf_counter()
    final_model.fit(development[feature_cols], development["label"])
    final_fit_seconds = time.perf_counter() - t0
    t0 = time.perf_counter()
    pred = final_model.predict(X["test"])
    infer_seconds = time.perf_counter() - t0

    result = pd.DataFrame([dict(
        classifier=winner,
        selected_variant_index=winner_idx,
        selection_validation_macro_f1=selected["validation_macro_f1"],
        test_accuracy=accuracy_score(y["test"], pred),
        test_balanced_accuracy=balanced_accuracy_score(y["test"], pred),
        test_macro_precision=precision_score(y["test"], pred, average="macro", zero_division=0),
        test_macro_recall=recall_score(y["test"], pred, average="macro", zero_division=0),
        test_macro_f1=f1_score(y["test"], pred, average="macro", zero_division=0),
        test_mcc=matthews_corrcoef(y["test"], pred),
        final_refit_seconds=final_fit_seconds,
        inference_seconds=infer_seconds,
    )])
    result.to_csv(RESULTS / "tables" / "selected_model_test_metrics.csv", index=False)
    joblib.dump(final_model, MODELS / "selected_model.joblib")
    (MODELS / "selected_model.json").write_text(json.dumps({
        "classifier": winner,
        "selected_variant_index": winner_idx,
        "selection_basis": "training fit and validation-only ranking; one test evaluation after refitting on train+validation",
        "feature_columns": feature_cols,
        "sampling_rate_hz_assumed": FS,
        "window_samples": WINDOW,
    }, indent=2), encoding="utf-8")

    report = classification_report(y["test"], pred,
                                   labels=CLASS_ORDER, output_dict=True, zero_division=0)
    pd.DataFrame(report).T.to_csv(RESULTS / "tables" / "selected_model_classification_report.csv")

    # Feature importance when available.
    cm = confusion_matrix(y["test"], pred, labels=CLASS_ORDER)
    plt.figure(figsize=(6.2, 5.2))
    sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", xticklabels=CLASS_ORDER, yticklabels=CLASS_ORDER)
    plt.xlabel("Predicted class")
    plt.ylabel("True class")
    plt.title(f"Selected {winner} - blocked test set")
    plt.tight_layout()
    plt.savefig(RESULTS / "confusion_matrices" / "selected_model.png", dpi=220)
    plt.close()

    pipe = final_model
    estimator = pipe.named_steps["model"]
    if hasattr(estimator, "feature_importances_"):
        imp = pd.DataFrame({"feature": feature_cols, "importance": estimator.feature_importances_}).sort_values("importance", ascending=False)
        imp.to_csv(RESULTS / "feature_importance" / "selected_model_feature_importance.csv", index=False)
        plt.figure(figsize=(8, 6))
        sns.barplot(data=imp.head(15), y="feature", x="importance", color="#2878B5")
        plt.title(f"Top feature importances - {winner}")
        plt.tight_layout()
        plt.savefig(RESULTS / "feature_importance" / "selected_model_feature_importance.png", dpi=220)
        plt.close()
    return result


def write_inventory():
    rows=[]
    for p in sorted(PROJECT.rglob("*")):
        if p.is_file() and ".venv" not in p.parts and not str(p).startswith(str(RESULTS)):
            rows.append(dict(path=str(p.relative_to(PROJECT)), size_bytes=p.stat().st_size,
                             sha256=sha256(p), suffix=p.suffix.lower()))
    pd.DataFrame(rows).to_csv(MANIFEST / "project_file_inventory.csv", index=False)


def main():
    ensure_dirs()
    write_inventory()
    runs = extract_current_runs()
    by_class, split_map = make_guarded_blocks(runs)
    df, feature_cols = build_feature_table(by_class, split_map)
    result = evaluate(df, feature_cols)
    print(result.to_string(index=False))


if __name__ == "__main__":
    main()
